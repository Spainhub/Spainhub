import io
import json
import uuid
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.db.session import set_rls_context
from tests.conftest import needs_db, setup_org
from tests.test_sync_db import FakeClient, run_all

pytestmark = needs_db
ROWS = json.loads((Path(__file__).parent / "fixtures" / "wb_real_sample.json").read_text(encoding="utf-8"))
SALES = [r for r in ROWS if r["sellerOperName"] == "Продажа"]
D1, D2 = "2026-07-01", "2026-07-11"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class WbFake(FakeClient):
    source_endpoint = "wb:sales-reports/detailed"


def xlsx(rows, headers=("subjectName", "nmId", "vendorCode", "costPerUnit", "validFrom")) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
async def world(admin, login_as, maker):
    tag = uuid.uuid4().hex[:6]
    a = await setup_org(admin, f"P-{tag}", [(f"adm-{tag}@example.com", "admin"), (f"mem-{tag}@example.com", "member"), (f"vw-{tag}@example.com", "viewer")])
    b = await setup_org(admin, f"Q-{tag}", [(f"b-{tag}@example.com", "admin")])
    w = {"a": a, "b": b, "admin": admin, "adm": await login_as(f"adm-{tag}@example.com"), "mem": await login_as(f"mem-{tag}@example.com"),
         "vw": await login_as(f"vw-{tag}@example.com"), "bad": await login_as(f"b-{tag}@example.com")}
    c, h = w["adm"]
    acc = (await c.post(f"/api/v1/orgs/{a['id']}/seller-accounts", headers=h, json={"name": "WB", "token": "x" * 30})).json()
    async with maker() as db:
        await set_rls_context(db, platform_admin=True)
        await db.execute(text("UPDATE sync_jobs SET status='cancelled' WHERE status IN ('queued','running')"))
        await db.commit()
    assert (await c.post(f"/api/v1/orgs/{a['id']}/seller-accounts/{acc['id']}/sync", headers=h, json={"months": 3})).status_code == 202
    await run_all(maker, WbFake([ROWS, []]))
    w["acc"] = acc["id"]
    return w


def url(w, path, org="a"):
    return f"/api/v1/orgs/{w[org]['id']}{path}"


async def set_regime(w, regime="usn_income", who="adm", **kw):
    c, h = w[who]
    return await c.put(url(w, "/settings"), headers=h, json={"tax_regime": regime, **kw})


async def confirm_manual(w):
    c, h = w["adm"]
    for cat in ("opex_other", "marketing_other"):
        r = await c.post(url(w, "/manual-entries"), headers=h, json={"entry_date": "2026-07-05", "category": cat, "amount": "0", "description": "подтверждено"})
        assert r.status_code == 201, r.text


async def upload_costs(w, rows, who="adm"):
    c, h = w[who]
    return await c.post(url(w, "/costs/import"), headers=h, files={"file": ("c.xlsx", xlsx(rows), XLSX)})


def line(rep, lid, col="total"):
    return next(l for l in rep["lines"] if l["id"] == lid)["values"][col]


async def report(w, **params):
    c, _ = w["adm"]
    q = {"date_from": D1, "date_to": D2, "granularity": "total", **params}
    r = await c.get(url(w, "/reports/pnl"), params=q)
    return r


# ----------------------------- настройки и права -----------------------------
async def test_report_requires_settings_and_settings_require_admin(world):
    r = await report(world)
    assert r.status_code == 409 and "учётная политика" in r.text
    assert (await set_regime(world, who="mem")).status_code == 403
    assert (await set_regime(world, who="vw")).status_code == 403
    assert (await set_regime(world, "bogus")).status_code == 422
    ok = await set_regime(world)
    assert ok.status_code == 200 and ok.json()["tax_regime"] == "usn_income" and D(ok.json()["usn_rate"]) == 6
    c, _ = world["vw"]
    assert (await c.get(url(world, "/settings"))).json()["tax_regime"] == "usn_income"


# ----------------------------- сквозной расчёт на реальных данных -----------------------------
async def test_end_to_end_usn_matches_independent_calculation(world):
    await set_regime(world)
    c, h = world["adm"]
    # шаблон содержит товары из данных
    t = await c.get(url(world, "/costs/template"))
    assert t.status_code == 200 and t.headers["content-type"].startswith(XLSX)
    ws = load_workbook(io.BytesIO(t.content)).active
    header = [x.value for x in ws[1]]
    assert header == ["subjectName", "nmId", "vendorCode", "costPerUnit", "validFrom"]
    nm_ids = {int(r[1].value) for r in ws.iter_rows(min_row=2)}
    assert nm_ids == {r["nmId"] for r in SALES}

    # без себестоимости валовая прибыль заблокирована, но выручка считается
    rep = (await report(world)).json()
    assert line(rep, "net_revenue")["status"] == "ok"
    assert line(rep, "cogs_goods")["status"] == "blocked" and line(rep, "gross_profit")["value"] is None
    assert line(rep, "gross_profit")["preliminary"] is not None

    # заполняем себестоимость: 1000 ₽ за единицу для всех
    up = await upload_costs(world, [(None, nm, None, "1000", "2026-01-01") for nm in sorted(nm_ids)])
    assert up.status_code == 200 and up.json()["imported"] == len(nm_ids)
    await confirm_manual(world)

    rep = (await report(world)).json()
    n = len(SALES)
    revenue = sum(D(str(r["retailPriceWithDisc"])) for r in SALES)
    commission = sum(D(str(r["retailPrice"])) * D(str(r["commissionPercent"])) / 100 for r in SALES)
    acquiring = sum(D(str(r["acquiringFee"])) for r in SALES)
    logistics = sum(D(str(r["deliveryService"])) for r in ROWS if r["sellerOperName"] == "Логистика")
    storage = sum(D(str(r["paidStorage"])) for r in ROWS if r["sellerOperName"] == "Хранение")
    accept = sum(D(str(r["paidAcceptance"])) for r in ROWS if r["sellerOperName"] == "Обработка товара")
    adv = sum(D(str(r["deduction"])) for r in ROWS if r["sellerOperName"] == "Удержание")
    cogs = D(1000) * sum(r["quantity"] for r in SALES)
    tol = D("0.02") * (n + len(ROWS))
    v = lambda lid: D(line(rep, lid)["value"])
    assert v("sales_gross") == revenue and v("net_revenue") == revenue
    assert abs(v("commission") - commission) <= tol and v("acquiring") == acquiring.quantize(D("0.01"))
    assert v("logistics") == logistics.quantize(D("0.01")) and v("storage") == storage.quantize(D("0.01"))
    assert v("acceptance") == accept and v("advertising_wb") == adv and v("cogs") == cogs
    assert abs(v("gross_profit") - (revenue - commission - acquiring - logistics - storage - accept - cogs)) <= tol
    assert abs(v("operating_profit") - (v("gross_profit") - adv)) <= D("0.01")
    assert v("tax_usn") == (revenue * D("0.06")).quantize(D("0.01"), rounding="ROUND_HALF_UP")
    assert abs(v("net_profit") - (v("operating_profit") - v("tax_usn"))) <= D("0.01")
    assert line(rep, "net_profit")["status"] in ("ok", "incomplete")
    # forPay — справка для сверки, не прибыль
    assert abs(D(line(rep, "wb_for_pay")["value"]) - (revenue - commission - acquiring)) <= tol
    assert rep["quality"]["transactions_by_status"]["review"] == 2        # возмещение издержек: ждёт политики
    assert "rebill_logistics_cost" in rep["quality"]["excluded_pending_amounts"]
    assert rep["coverage"]["days_with_data"] > 0


async def test_month_and_day_granularity_and_validation(world):
    await set_regime(world)
    rep = (await report(world, granularity="month")).json()
    assert rep["periods"] == ["2026-07", "total"]
    assert line(rep, "sales_gross", "2026-07")["value"] == line(rep, "sales_gross", "total")["value"]
    rep = (await report(world, granularity="day")).json()
    assert len(rep["periods"]) == 12 and sum(D(line(rep, "sales_gross", p)["value"] or 0) for p in rep["periods"][:-1]) == D(line(rep, "sales_gross")["value"])
    assert (await report(world, date_from="2026-01-01", date_to="2026-12-31", granularity="day")).status_code == 422
    assert (await report(world, date_from="2026-07-11", date_to="2026-07-01")).status_code == 422
    # период без импорта — «нет данных», а не нули
    empty = (await report(world, date_from="2025-01-01", date_to="2025-01-31")).json()
    assert line(empty, "sales_gross")["status"] == "missing" and line(empty, "net_profit")["value"] is None


async def test_osno_report_has_vat_and_profit_tax(world):
    await set_regime(world, "osno", profit_tax_rate="20", vat_sales_rate="22", vat_services_rate="22")
    rep = (await report(world)).json()
    ids = {l["id"] for l in rep["lines"]}
    assert {"vat_output", "tax_profit", "vat_payable"} <= ids and "tax_usn" not in ids
    gross, vat, net = (D(line(rep, k)["value"]) for k in ("sales_gross", "vat_output", "net_revenue"))
    assert vat == (gross * 22 / 122).quantize(D("0.01")) and net == gross - vat


# ----------------------------- импорт себестоимости -----------------------------
async def test_cost_import_is_all_or_nothing_and_validates(world):
    c, h = world["adm"]
    nm = SALES[0]["nmId"]
    bad = await upload_costs(world, [(None, nm, None, "100", "2026-01-01"), (None, "abc", None, "5", None), (None, nm + 1, None, "-3", None),
                                     (None, nm + 2, None, "1,5", "31-12-2026x")])
    assert bad.status_code == 422
    fields = {(e["row"], e["field"]) for e in bad.json()["detail"]["errors"]}
    assert {(3, "nmId"), (4, "costPerUnit"), (5, "validFrom")} <= fields
    assert (await c.get(url(world, "/costs"))).json() == []                  # ничего не загружено
    dup = await upload_costs(world, [(None, nm, None, "1", "2026-01-01"), (None, nm, None, "2", "2026-01-01")])
    assert dup.status_code == 422 and "Повтор" in dup.text


async def test_cost_import_rejects_junk_and_oversize(world):
    c, h = world["adm"]
    for payload in (b"not a zip", b"PK\x03\x04garbage", b"x" * (2 * 1024 * 1024 + 10)):
        r = await c.post(url(world, "/costs/import"), headers=h, files={"file": ("c.xlsx", payload, XLSX)})
        assert r.status_code == 422, r.status_code
    r = await c.post(url(world, "/costs/import"), headers=h, files={"file": ("c.xlsx", xlsx([], headers=("a", "b")), XLSX)})
    assert r.status_code == 422 and "nmId" in r.text


async def test_blank_cost_rows_skipped_and_cost_changes_by_date(world):
    await set_regime(world)
    nm = SALES[0]["nmId"]
    r = await upload_costs(world, [(None, nm, None, "100", None), (None, nm, None, "300", "2026-07-06"), (None, 999, "ART", None, None)])
    assert r.json() == {"imported": 2, "skipped_blank": 1}
    c, _ = world["adm"]
    items = (await c.get(url(world, "/reports/pnl/drilldown"), params={"line": "cogs_goods", "date_from": D1, "date_to": D2})).json()["items"]
    mine = [i for i in items if i["nm_id"] == nm]
    assert mine and all(D(i["cost_per_unit"]) == (D(300) if i["entry_date"] >= "2026-07-06" else D(100)) for i in mine)
    # повторная загрузка той же даты обновляет, а не дублирует
    await upload_costs(world, [(None, nm, None, "150", None)])
    costs = [x for x in (await c.get(url(world, "/costs"))).json() if x["nm_id"] == nm]
    assert sorted(D(x["cost_per_unit"]) for x in costs) == [D(150), D(300)]


async def test_viewer_cannot_import_costs_or_write_entries(world):
    assert (await upload_costs(world, [(None, 1, None, "1", None)], who="vw")).status_code == 403
    c, h = world["vw"]
    r = await c.post(url(world, "/manual-entries"), headers=h, json={"entry_date": "2026-07-01", "category": "opex_other", "amount": "1"})
    assert r.status_code == 403
    assert (await c.get(url(world, "/manual-entries"))).status_code == 200


# ----------------------------- ручные записи -----------------------------
async def test_manual_entries_crud_validation_and_report(world):
    await set_regime(world)
    c, h = world["mem"]
    ok = await c.post(url(world, "/manual-entries"), headers=h, json={"entry_date": "2026-07-03", "category": "opex_salaries", "amount": "1500.50", "description": "ФОТ"})
    assert ok.status_code == 201 and ok.json()["created_by"] is not None
    eid = ok.json()["id"]
    for bad in ({"category": "nonsense"}, {"amount": "-1"}, {"currency": "rub"}, {"entry_date": "not-a-date"}):
        payload = {"entry_date": "2026-07-03", "category": "opex_other", "amount": "1", **bad}
        assert (await c.post(url(world, "/manual-entries"), headers=h, json=payload)).status_code == 422, bad
    assert (await c.post(url(world, "/manual-entries"), headers={}, json={"entry_date": "2026-07-03", "category": "opex_other", "amount": "1"})).status_code == 403  # CSRF
    rep = (await report(world)).json()
    assert line(rep, "opex_salaries")["value"] == "1500.50" and line(rep, "opex")["status"] != "missing"
    up = await c.put(url(world, f"/manual-entries/{eid}"), headers=h, json={"entry_date": "2026-07-03", "category": "opex_salaries", "amount": "2000"})
    assert up.status_code == 200 and D(up.json()["amount"]) == 2000
    dr = (await c.get(url(world, "/reports/pnl/drilldown"), params={"line": "opex_salaries", "date_from": D1, "date_to": D2})).json()
    assert dr["kind"] == "manual" and dr["items"][0]["id"] == eid
    assert (await c.delete(url(world, f"/manual-entries/{eid}"), headers=h)).status_code == 204
    assert (await c.delete(url(world, f"/manual-entries/{eid}"), headers=h)).status_code == 404


async def test_drilldown_import_lines_and_rejects_formula_lines(world):
    await set_regime(world)
    c, _ = world["adm"]
    dr = (await c.get(url(world, "/reports/pnl/drilldown"), params={"line": "commission", "date_from": D1, "date_to": D2})).json()
    assert dr["kind"] == "import" and len(dr["items"]) == len(SALES) and {"source_id", "amount", "transaction_id"} <= set(dr["items"][0])
    assert (await c.get(url(world, "/reports/pnl/drilldown"), params={"line": "gross_profit", "date_from": D1, "date_to": D2})).status_code == 422
    assert (await c.get(url(world, "/reports/pnl/drilldown"), params={"line": "nope", "date_from": D1, "date_to": D2})).status_code == 404


# ----------------------------- изоляция организаций -----------------------------
async def test_other_organization_cannot_reach_pnl_data(world):
    await set_regime(world)
    await confirm_manual(world)
    c, h = world["bad"]
    a = world["a"]["id"]
    for path in ("/settings", "/costs", "/manual-entries", "/reports/pnl?date_from=2026-07-01&date_to=2026-07-11", "/costs/template"):
        assert (await c.get(f"/api/v1/orgs/{a}{path}")).status_code == 404, path
    assert (await c.put(f"/api/v1/orgs/{a}/settings", headers=h, json={"tax_regime": "osno"})).status_code == 404
    # свой P&L не содержит чужих данных
    await c.put(url(world, "/settings", "b"), headers=h, json={"tax_regime": "usn_income"})
    rep = (await c.get(url(world, "/reports/pnl", "b"), params={"date_from": D1, "date_to": D2, "granularity": "total"})).json()
    assert line(rep, "sales_gross")["status"] == "missing"
    # чужой аккаунт и чужая ручная запись
    r = await c.get(url(world, "/reports/pnl", "b"), params={"date_from": D1, "date_to": D2, "account_id": world["acc"]})
    assert r.status_code == 404
    r = await c.post(url(world, "/manual-entries", "b"), headers=h, json={"entry_date": D1, "category": "opex_other", "amount": "1", "seller_account_id": world["acc"]})
    assert r.status_code == 404
    mine = (await world["adm"][0].get(url(world, "/manual-entries"))).json()[0]["id"]
    assert (await c.delete(url(world, f"/manual-entries/{mine}", "b"), headers=h)).status_code == 404
    assert (await world["adm"][0].get(url(world, "/manual-entries"))).json()


async def test_rls_on_pnl_tables(world, maker):
    await set_regime(world)
    await confirm_manual(world)
    await upload_costs(world, [(None, SALES[0]["nmId"], None, "1", None)])
    async with maker() as db:
        for t in ("org_settings", "product_costs", "manual_entries"):
            assert (await db.execute(text(f"SELECT count(*) FROM {t}"))).scalar_one() == 0
        await set_rls_context(db, org_id=uuid.uuid4())
        for t in ("org_settings", "product_costs", "manual_entries"):
            assert (await db.execute(text(f"SELECT count(*) FROM {t}"))).scalar_one() == 0
        await set_rls_context(db, org_id=world["a"]["id"])
        for t in ("org_settings", "product_costs", "manual_entries"):
            assert (await db.execute(text(f"SELECT count(*) FROM {t}"))).scalar_one() >= 1
