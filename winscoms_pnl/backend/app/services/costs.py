"""Себестоимость: Excel-шаблон (скачать → заполнить → загрузить)."""
import io
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import ProductCost

HEADERS = ["subjectName", "nmId", "vendorCode", "costPerUnit", "validFrom"]
SHEET = "Себестоимость"
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 10000
DEFAULT_VALID_FROM = date(1970, 1, 1)
MAX_COST = Decimal("1000000000")


class CostImportError(Exception):
    def __init__(self, errors: list[dict]) -> None:
        super().__init__("Ошибки в файле")
        self.errors = errors


async def products_for_template(db: AsyncSession, org_id: uuid.UUID) -> list[tuple[str | None, int, str | None]]:
    """Товары, встречавшиеся в импортированных продажах: (subjectName, nmId, vendorCode)."""
    rows = await db.execute(text("""
        SELECT DISTINCT ON (ft.nm_id) rr.payload->>'subjectName', ft.nm_id, ft.vendor_code
        FROM financial_transactions ft JOIN raw_records rr ON rr.id = ft.raw_record_id
        WHERE ft.organization_id = :org AND ft.category = 'sale' AND ft.nm_id IS NOT NULL
        ORDER BY ft.nm_id, ft.entry_date DESC"""), {"org": org_id})
    return [(r[0], r[1], r[2]) for r in rows]


def build_template(products: list[tuple[str | None, int, str | None]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for subject, nm_id, vendor in products:
        ws.append([subject, nm_id, vendor, None, None])
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "@"              # артикул — текст (Excel иначе съедает ведущие нули)
        row[3].number_format = "0.00"
        row[4].number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDE", (28, 14, 22, 14, 14)):
        ws.column_dimensions[col].width = w
    help_ws = wb.create_sheet("Инструкция")
    for line in (
        "Заполните costPerUnit — себестоимость ОДНОЙ единицы товара (для режима ОСНО — без НДС).",
        "validFrom — дата, с которой действует себестоимость (формат ГГГГ-ММ-ДД). Пусто = действует с начала.",
        "Чтобы изменить себестоимость с определённой даты, добавьте новую строку с тем же nmId и новой validFrom.",
        "Строки с пустым costPerUnit пропускаются. nmId обязателен. subjectName и vendorCode — справочные.",
        "Если загрузка содержит хотя бы одну ошибку, не загружается ничего — исправьте файл и повторите.",
    ):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _err(row: int, field: str, msg: str) -> dict:
    return {"row": row, "field": field, "error": msg}


def _int(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def _decimal(v) -> Decimal | None:
    if isinstance(v, bool):
        return None
    try:
        d = Decimal(str(v).strip().replace(",", ".").replace(" ", "").replace(" ", ""))
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def parse_workbook(content: bytes) -> tuple[list[dict], int]:
    """Возвращает (строки, число пропущенных пустых). Бросает CostImportError со списком ошибок."""
    if len(content) > MAX_FILE_BYTES:
        raise CostImportError([_err(0, "file", f"Файл больше {MAX_FILE_BYTES // 1024 // 1024} МБ")])
    if not content.startswith(b"PK"):
        raise CostImportError([_err(0, "file", "Это не файл .xlsx")])
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise CostImportError([_err(0, "file", "Файл слишком большой после распаковки")])
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except CostImportError:
        raise
    except Exception:
        raise CostImportError([_err(0, "file", "Не удалось прочитать файл .xlsx")])
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        raise CostImportError([_err(1, "header", "Пустой файл")])
    cols = {str(h).strip().lower(): i for i, h in enumerate(header) if h is not None}
    idx = {h: cols.get(h.lower()) for h in HEADERS}
    missing = [h for h in ("nmId", "costPerUnit") if idx[h] is None]
    if missing:
        raise CostImportError([_err(1, "header", f"Нет обязательных колонок: {', '.join(missing)}")])

    def cell(row, name):
        i = idx[name]
        return row[i] if i is not None and i < len(row) else None

    out, errors, skipped, seen = [], [], 0, set()
    for n, row in enumerate(it, start=2):
        if n - 1 > MAX_ROWS:
            errors.append(_err(n, "file", f"Больше {MAX_ROWS} строк"))
            break
        cost_raw = cell(row, "costPerUnit")
        if cost_raw is None or str(cost_raw).strip() == "":
            skipped += 1
            continue
        nm_id = _int(cell(row, "nmId"))
        if nm_id is None or nm_id <= 0:
            errors.append(_err(n, "nmId", "Нужен числовой nmId"))
        cost = _decimal(cost_raw)
        if cost is None or cost < 0 or cost > MAX_COST:
            errors.append(_err(n, "costPerUnit", "Нужно число от 0 (допустимы десятичные, не более 6 знаков)"))
        elif -cost.as_tuple().exponent > 6:
            errors.append(_err(n, "costPerUnit", "Не более 6 знаков после запятой"))
        vf_raw = cell(row, "validFrom")
        valid_from = DEFAULT_VALID_FROM
        if vf_raw is not None and str(vf_raw).strip() != "":
            valid_from = _date(vf_raw)
            if valid_from is None:
                errors.append(_err(n, "validFrom", "Дата в формате ГГГГ-ММ-ДД"))
        vendor = cell(row, "vendorCode")
        if nm_id and cost is not None and valid_from:
            key = (nm_id, valid_from)
            if key in seen:
                errors.append(_err(n, "nmId", "Повтор nmId с той же validFrom"))
            seen.add(key)
            out.append({"nm_id": nm_id, "cost_per_unit": cost, "valid_from": valid_from,
                        "vendor_code": None if vendor is None else str(vendor)})
    if errors:
        raise CostImportError(errors[:100])
    return out, skipped


async def import_costs(db: AsyncSession, org_id: uuid.UUID, user_id, rows: list[dict], currency: str) -> int:
    for i in range(0, len(rows), 2000):
        chunk = [{**r, "organization_id": org_id, "currency": currency, "created_by": user_id} for r in rows[i : i + 2000]]
        stmt = pg_insert(ProductCost).values(chunk)
        stmt = stmt.on_conflict_do_update(
            index_elements=["organization_id", "nm_id", "currency", "valid_from"],
            set_={"cost_per_unit": stmt.excluded.cost_per_unit, "vendor_code": stmt.excluded.vendor_code, "created_by": user_id},
        )
        await db.execute(stmt)
    return len(rows)
